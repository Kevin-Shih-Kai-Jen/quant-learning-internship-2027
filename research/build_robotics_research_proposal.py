from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "機器人與自動化股票研究企劃書.docx"
SKILL = Path("/Users/coolguy/.codex/plugins/cache/openai-primary-runtime/documents/26.819.11345/skills/documents")
sys.path.insert(0, str(SKILL / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F6B76"
LIGHT = "F4F6F9"
BLUE_LIGHT = "E8EEF5"
GOLD_LIGHT = "FFF8E8"
GOLD = "9A6A00"
RED_LIGHT = "FDECEC"
RED = "9B1C1C"
GREEN_LIGHT = "EAF5EF"
GREEN = "1E6A43"
WHITE = "FFFFFF"
BLACK = "111111"
GRID = "C9D2DC"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=None, bold=None, italic=None, color=BLACK,
                 latin="Heiti TC", east_asia="Heiti TC"):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def set_style_font(style, size, color=BLACK, bold=None, east_asia="Heiti TC"):
    style.font.name = "Heiti TC"
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Heiti TC")
    rfonts.set(qn("w:hAnsi"), "Heiti TC")
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        if edge in edges:
            edge_data = edges.get(edge)
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(color)
    rpr.append(sz)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(rpr)
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)
    run = paragraph.add_run(" 頁")
    set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Heiti TC")
    rfonts.set(qn("w:hAnsi"), "Heiti TC")
    rfonts.set(qn("w:eastAsia"), "Heiti TC")
    rpr.append(rfonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    new_run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = doc.styles["Title"]
    set_style_font(title, 29, NAVY, True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.05

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, 13.2, MUTED, False)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        set_style_font(style, size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = doc.styles[name]
        set_style_font(style, 11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    if "Table Text" not in [s.name for s in doc.styles]:
        table_style = doc.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_style = doc.styles["Table Text"]
    set_style_font(table_style, 9.5)
    table_style.paragraph_format.space_before = Pt(0)
    table_style.paragraph_format.space_after = Pt(2)
    table_style.paragraph_format.line_spacing = 1.08

    if "Source Text" not in [s.name for s in doc.styles]:
        src_style = doc.styles.add_style("Source Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        src_style = doc.styles["Source Text"]
    set_style_font(src_style, 9, MUTED)
    src_style.paragraph_format.space_before = Pt(4)
    src_style.paragraph_format.space_after = Pt(4)
    src_style.paragraph_format.line_spacing = 1.1

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    run = hp.add_run("機器人與自動化股票研究企劃書")
    set_run_font(run, size=9, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_field(fp)
    return doc


def add_para(doc, text="", bold=False, italic=False, color=BLACK, size=None,
             align=None, before=0, after=8, style=None, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
    r = p.add_run(text)
    set_run_font(r)
    return p


def new_number_sequence(doc):
    """Create an independent decimal list that restarts at 1."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = numbering.num_having_numId(style_num_id)
    num = numbering.add_num(base_num.abstractNumId.val)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return int(num.numId)


def add_number(doc, text, num_id=None):
    p = doc.add_paragraph(style="List Number")
    if num_id is not None:
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num_id
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_callout(doc, label, text, fill=BLUE_LIGHT, accent=BLUE, text_color=BLACK):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell,
                    start={"val": "single", "sz": "24", "color": accent},
                    top={"val": "single", "sz": "4", "color": fill},
                    bottom={"val": "single", "sz": "4", "color": fill},
                    end={"val": "single", "sz": "4", "color": fill})
    p = cell.paragraphs[0]
    p.style = "Table Text"
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label + "  ")
    set_run_font(r, size=10.5, bold=True, color=accent)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=text_color)
    apply_table_geometry(table, [9360], indent_dxa=160,
                         cell_margins_dxa={"top": 130, "bottom": 130, "start": 160, "end": 160})
    add_para(doc, "", after=2)
    return table


def add_table(doc, headers, rows, widths, header_fill=LIGHT, font_size=9.3,
              aligns=None, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.style = "Table Text"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(str(text))
        set_run_font(r, size=font_size, bold=True, color=NAVY)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = "Table Text"
            p.alignment = (aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, bold=(first_col_bold and i == 0), color=BLACK)
    apply_table_geometry(table, widths, indent_dxa=120,
                         cell_margins_dxa={"top": 90, "bottom": 90, "start": 120, "end": 120})
    add_para(doc, "", after=2)
    return table


def add_source_note(doc, text):
    p = doc.add_paragraph(style="Source Text")
    r = p.add_run(text)
    set_run_font(r, size=9, color=MUTED)
    return p


def add_page_break(doc):
    doc.add_page_break()


def add_cover(doc):
    add_para(doc, "研究與決策工作手冊", size=10, bold=True, color=BLUE,
             before=16, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("機器人與自動化股票研究企劃書")
    set_run_font(r, size=29, bold=True, color=NAVY)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("第一次研究會議、四週研究流程與一年目標價框架")
    set_run_font(r, size=13.2, color=MUTED)

    metrics = [
        ("投資期", "1年"),
        ("最終目標", "選出2-3家公司"),
        ("風險界線", "單檔最多約-30%"),
        ("首次會議", "90分鐘"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, (label, value) in enumerate(metrics):
        cell = table.cell(0, i)
        set_cell_shading(cell, GOLD_LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label + "\n")
        set_run_font(r, size=8.5, bold=True, color=GOLD)
        r = p.add_run(value)
        set_run_font(r, size=10.2, bold=True, color=NAVY)
    apply_table_geometry(table, [2340, 2340, 2340, 2340], indent_dxa=80,
                         cell_margins_dxa={"top": 150, "bottom": 150, "start": 80, "end": 80})

    add_para(doc, "", after=10)
    add_callout(
        doc,
        "核心任務",
        "在未來12個月內，找出能把技術優勢轉成訂單、毛利率與EPS的公司，並以悲觀、基準、樂觀三種情境估值，判斷預期報酬是否值得承擔風險。",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )
    add_para(doc, "本企劃書不是明牌清單，而是一套共同研究制度。第一次會議不需要決定買進，而要完成研究範圍、證據標準、分工與下一次交付物。",
             size=11.2, color=NAVY, before=8, after=16)

    meta = add_table(doc,
        ["項目", "內容"],
        [
            ["參與者", "投資／資料主責（你）＋技術／產業主責（機械系朋友）"],
            ["研究市場", "台股與一般美股券商可交易之美國上市公司"],
            ["版本日期", "2026年8月25日"],
            ["文件狀態", "第一次會議前閱讀版；所有估值均須以最新財報與股價更新"],
        ],
        [1800, 7560],
        font_size=9.6,
        first_col_bold=True,
    )
    return meta


def add_toc_and_summary(doc):
    add_page_break(doc)
    doc.add_heading("閱讀方式與企劃摘要", level=1)
    add_para(doc, "建議會前先花30分鐘閱讀第1至第6章；開會時直接使用第4章議程；會後依第10章時程執行。附錄提供可重複使用的研究卡、估值表與會議紀錄格式。")

    doc.add_heading("企劃書目錄", level=2)
    toc_num_id = new_number_sequence(doc)
    for item in [
        "研究目標與投資邊界",
        "研究範圍與公司分類",
        "一年投資期的判斷原則",
        "第一次會議：90分鐘完整議程",
        "兩人分工與交叉審查",
        "技術與競爭力研究方法",
        "財務預測與一年目標價模型",
        "風險控制、停損與部位管理",
        "公司研究優先順序",
        "四週研究計畫與交付物",
        "最終選股與投資決策規則",
        "附錄A-C：研究卡、估值表與會議紀錄",
        "資料來源與使用限制",
    ]:
        add_number(doc, item, toc_num_id)

    doc.add_heading("第一次會議結束時，必須產出五件事", level=2)
    for text in [
        "確定研究母體：哪些是直接機器人公司、零組件公司、工業自動化公司，以及哪些應移出同業比較。",
        "選出第一輪深度研究的4-5家公司，而不是一次研究所有題材股。",
        "建立共同證據標準：什麼算量產、什麼只是展示或市場敘事。",
        "完成兩人分工、交付格式與截止時間。",
        "定下第二次會議要回答的核心問題，不在第一次會議倉促買進。",
    ]:
        add_bullet(doc, text)

    add_callout(doc, "一句話原則", "機器人產品存在，不代表機器人營收已經存在；營收成長，也不代表目前估值值得買進。", fill=GOLD_LIGHT, accent=GOLD)


def section_1(doc):
    doc.add_heading("1. 研究目標與投資邊界", level=1)
    doc.add_heading("1.1 已確定的投資條件", level=2)
    add_table(doc,
        ["條件", "目前設定", "對研究方法的影響"],
        [
            ["投資期", "1年", "以未來四季訂單、營收、毛利率和EPS為主；三年以上題材只影響估值溢價。"],
            ["持股數", "2-3家公司", "研究必須能比較和淘汰，不以蒐集大量概念股為目標。"],
            ["風險偏好", "可接受未成熟與高波動", "可納入SYM、早期機器人零組件等高風險候選，但必須要求更高上漲空間。"],
            ["價格停損", "單檔約-30%", "需同時搭配基本面停損與部位上限；停損價不能取代研究。"],
            ["決策方式", "研究後再買", "第一次會議只建制度與任務，不以當天股價波動改變研究順序。"],
        ],
        [1500, 1800, 6060],
        first_col_bold=True,
    )

    doc.add_heading("1.2 研究要回答的總問題", level=2)
    add_callout(doc, "總問題", "哪2-3家公司最有機會在未來12個月內，同時出現可驗證的營運改善、足以防禦競爭的技術或商業優勢，以及尚未被股價完全反映的報酬空間？")

    doc.add_heading("1.3 不在本輪研究範圍內", level=2)
    for text in [
        "只因媒體稱為機器人概念股，就假設公司會受惠。",
        "用三至五年後的人形機器人市場規模，直接推算一年目標價。",
        "只看營收成長，不檢查毛利率、現金流、股本稀釋與客戶集中度。",
        "用單一券商目標價代替自己的EPS與估值假設。",
        "以-30%價格停損作為唯一風險管理方法。",
    ]:
        add_bullet(doc, text)


def section_2(doc):
    doc.add_heading("2. 研究範圍與公司分類", level=1)
    add_para(doc, "研究前先分清楚公司在價值鏈中的位置。不同位置的營收模式、競爭者、估值方式和風險完全不同，不宜使用同一套機器人敘事。")
    add_table(doc,
        ["公司", "分類", "直接度", "一年內最重要的驗證變數"],
        [
            ["直得 1597", "精密運動元件＋微型機械手臂／控制", "中高", "微型手臂是否從產品展示進入客戶量產；線性滑軌與半導體設備需求能否支撐EPS。"],
            ["大銀微系統 4576", "馬達、驅動器、定位平台與單軸機器人", "高（零組件）", "擴產、精密定位平台與自動化元件訂單能否持續轉成毛利與EPS。"],
            ["波若威 3163", "光通訊元件與模組", "低", "CPO／800G／1.6T量產節奏；應獨立歸類為AI光通訊，不與機器人零組件直接比較。"],
            ["ISRG", "手術機器人系統＋耗材與服務", "高", "手術量、裝機量與每台系統的持續性收入；競爭與醫療監管。"],
            ["TER", "協作機器人／AMR＋半導體測試", "中", "機器人部門復甦，但整體一年獲利主要仍受AI半導體測試影響。"],
            ["SYM", "AI倉儲機器人與軟體平台", "高", "部署速度、系統毛利與Walmart客戶集中風險。"],
            ["ROK", "工業自動化、控制與服務", "中低", "作為成熟自動化公司與景氣循環估值基準，不是純機器人股。"],
        ],
        [1300, 2100, 900, 5060],
        font_size=8.8,
        first_col_bold=True,
    )
    add_source_note(doc, "產品與公司分類來源：公司官網與最新公開財報，詳見资料來源 [S1]-[S3]、[S9]-[S13]。")

    doc.add_heading("2.1 建議的第一輪研究母體", level=2)
    add_para(doc, "核心研究池先限制為5家公司：大銀、直得、ISRG、TER、SYM。波若威另建AI光通訊研究表；ROK只作估值與商業模式基準。這樣可以保留台灣零組件、美國成熟商業模式與高成長純題材三種角度。")

    doc.add_heading("2.2 研究層級", level=2)
    for text in [
        "核心候選：已存在營收與獲利，且未來四季有明確催化劑。",
        "高風險候選：商業化仍在擴張期，但若成功，營收與估值可能大幅變化。",
        "比較基準：未必買進，但用來判斷毛利、客戶黏著度、成長品質與合理本益比。",
        "觀察名單：題材吸引人，但目前無法把技術敘事連到未來四季財務數字。",
    ]:
        add_bullet(doc, text)


def section_3(doc):
    doc.add_heading("3. 一年投資期的判斷原則", level=1)
    add_callout(doc, "時間軸切割", "一年目標價只能主要反映未來12個月EPS或營收。人形機器人三至五年的選擇權可以給予估值溢價，但不能當作基準情境的主要獲利。", fill=GOLD_LIGHT, accent=GOLD)

    doc.add_heading("3.1 五道投資關卡", level=2)
    gates = [
        ("關卡1：現在靠什麼賺錢", "先理解既有產品、客戶、毛利與現金流，避免只研究新故事。"),
        ("關卡2：未來四季如何成長", "催化劑必須能寫成出貨、ASP、產能利用率、產品組合或費用率的變化。"),
        ("關卡3：成長能否持續", "檢查客戶驗證、切換成本、可靠度、售後服務和軟硬體整合。"),
        ("關卡4：市場已經期待多少", "用反推EPS與反推營收判斷股價是否已經包含樂觀情境。"),
        ("關卡5：錯了如何退出", "預先寫下價格停損、基本面停損、檢查日期與最大部位。"),
    ]
    for title, detail in gates:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(title + "：")
        set_run_font(r, bold=True, color=DARK_BLUE)
        r = p.add_run(detail)
        set_run_font(r)

    doc.add_heading("3.2 合格投資論點的句型", level=2)
    add_callout(doc, "投資論點模板", "因為［客戶／產業需求］在未來四季推動［產品出貨、價格或產品組合］，公司營收與EPS可望達到［數字］；即使估值從［目前倍數］回到［合理倍數］，基準目標價仍有［至少25%-30%］空間。若［否決條件］發生，論點失效。")

    doc.add_heading("3.3 題材與證據的分界", level=2)
    add_table(doc,
        ["層級", "可接受的證據", "對估值的用途"],
        [
            ["敘事", "展覽、研發、合作意向、媒體訪問", "只能列為觀察，不直接計入基準營收。"],
            ["送樣／驗證", "客戶測試、認證進度、試產", "可放入樂觀情境並設定成功機率。"],
            ["小量出貨", "可辨識營收、交期、重複訂單", "部分計入基準情境，但要確認毛利。"],
            ["量產", "明確訂單、產能利用率、營收占比、現金流", "可成為基準EPS的主要驅動。"],
            ["持續收入", "耗材、維護、軟體或長期服務", "可支持較高且較穩定的估值倍數。"],
        ],
        [1200, 3900, 4260],
        first_col_bold=True,
    )


def section_4(doc):
    add_page_break(doc)
    doc.add_heading("4. 第一次會議：90分鐘完整議程", level=1)
    add_callout(doc, "會議目的", "不決定今天要買哪一檔；完成共同語言、研究母體、證據標準、分工與下一次交付物。")

    doc.add_heading("4.1 會前準備（各30-45分鐘）", level=2)
    add_table(doc,
        ["角色", "會前準備", "帶進會議的資料"],
        [
            ["你：投資／資料", "整理每家公司最新股價、近四季EPS、毛利率、營收成長、當前估值與重要財報日期。", "一張初步財務表；不需要先下結論。"],
            ["朋友：技術／產業", "查看直得與大銀產品頁，畫出機器人運動鏈；列出主要中國、美國、日本／歐洲競爭者。", "一張產品位置圖＋最想驗證的5個技術問題。"],
            ["共同", "各自選一間最看好和一間最懷疑的公司，寫下理由。", "會議開始時交換，不先說服對方。"],
        ],
        [1600, 5000, 2760],
        first_col_bold=True,
    )

    doc.add_heading("4.2 逐分鐘議程", level=2)
    add_table(doc,
        ["時間", "主題", "主持", "要回答的問題", "輸出"],
        [
            ["0-10", "目標與風險", "你", "一年後希望得到什麼？高波動可接受，但單檔與組合最大損失是多少？", "投資邊界確認"],
            ["10-22", "產業鏈地圖", "朋友", "每家公司位於本體、傳動、馬達／驅動、控制、視覺或應用哪一層？", "價值鏈草圖"],
            ["22-35", "公司分類", "共同", "波若威是否應移出機器人同業？哪些公司只是比較基準？", "第一輪5家公司"],
            ["35-50", "技術持續性", "朋友", "產品成熟度、客戶驗證、切換成本與中美自研風險如何？", "技術問題清單"],
            ["50-65", "財務與估值", "你", "股價目前要求多高的EPS？一年催化劑能否支持？", "估值假設清單"],
            ["65-75", "反方挑戰", "共同", "對方最看好的公司，有哪三個理由可能失敗？", "每家公司否決條件"],
            ["75-85", "分工與時程", "你", "誰負責哪些公司、資料格式、完成日期？", "任務表"],
            ["85-90", "確認下次會議", "共同", "下次要做淘汰、估值，還是買進決策？", "下次日期與議程"],
        ],
        [700, 1350, 900, 4310, 2100],
        font_size=8.5,
        first_col_bold=True,
    )

    doc.add_heading("4.3 建議開場白", level=2)
    add_callout(doc, "可直接照讀", "我們今天不是來挑明牌，而是要分辨哪些公司真的靠機器人或自動化賺錢、哪些只是題材；並決定用什麼證據和數字，讓我們在下一次會議能合理地淘汰公司。", fill=GREEN_LIGHT, accent=GREEN)

    doc.add_heading("4.4 會議規則", level=2)
    for text in [
        "先分類，再比較；不同商業模式不使用同一組估值倍數。",
        "每一項重要說法標記為：事實、公司說法、第三方估計、或自己的推論。",
        "不知道的問題直接列入待辦，不用在會議中爭論到有答案。",
        "任何人提出買進時，必須同時說出悲觀情境與否決條件。",
        "會議結束前，每項任務都要有負責人、期限與交付格式。",
    ]:
        add_bullet(doc, text)


def section_5(doc):
    doc.add_heading("5. 兩人分工與交叉審查", level=1)
    add_table(doc,
        ["工作流", "技術／產業主責（朋友）", "投資／資料主責（你）", "共同審查"],
        [
            ["產品", "拆解零件、規格、應用與成熟度。", "確認產品是否揭露營收或訂單。", "技術優勢能否轉成收入？"],
            ["競爭", "比較精度、壽命、可靠度、量產與認證。", "比較毛利率、價格壓力與市占證據。", "客戶為何不能換供應商？"],
            ["成長", "判斷客戶採用機率與所需時間。", "建立出貨、ASP、毛利、EPS模型。", "時間是否落在一年投資期內？"],
            ["風險", "找出中國低價、美國自研與技術替代。", "找出估值、客戶集中、稀釋與現金流風險。", "定義基本面停損。"],
            ["結論", "給技術分數與最大疑點。", "給目標價區間與預期報酬。", "決定淘汰、觀察或進入投資決策。"],
        ],
        [1300, 2650, 2650, 2760],
        font_size=8.8,
        first_col_bold=True,
    )

    doc.add_heading("5.1 建議的公司分配", level=2)
    for text in [
        "朋友主責：直得、大銀的技術與競爭比較；TER／Universal Robots作海外技術參照。",
        "你主責：所有公司的財務共同比較；先完成大銀的一年EPS與估值模型。",
        "高風險專案：你研究SYM的客戶集中與財務；朋友研究其倉儲系統的技術與競爭者。",
        "波若威：你獨立建立AI光通訊研究表，除非找到直接機器人營收證據，否則不放進機器人排名。",
    ]:
        add_bullet(doc, text)

    doc.add_heading("5.2 交叉審查問題", level=2)
    add_table(doc,
        ["朋友挑戰你的模型", "你挑戰朋友的技術結論"],
        [
            ["EPS成長來自真實產能與訂單，還是只把產業成長率套上去？", "規格較高，客戶是否真的願意付更高價格？"],
            ["毛利率改善需要哪些技術或產品組合條件？", "競爭者是否已有足夠可靠、但更便宜的替代品？"],
            ["本益比溢價是因為持續收入，還是市場情緒？", "技術驗證需要多久，能否在四季內影響財報？"],
            ["失敗時EPS和合理倍數會同時下修多少？", "公司說量產，究竟是小量、客戶試產還是規模出貨？"],
        ],
        [4680, 4680],
        font_size=9.2,
    )


def section_6(doc):
    doc.add_heading("6. 技術與競爭力研究方法", level=1)
    doc.add_heading("6.1 產品成熟度階梯", level=2)
    add_table(doc,
        ["階段", "可觀察證據", "研究處理方式"],
        [
            ["概念／展示", "展覽樣機、研發新聞、規格表", "只列為長期選擇權，不放入基準營收。"],
            ["送樣", "客戶測試、樣品收入、合作開發", "估計驗證成功率與時間，放入樂觀情境。"],
            ["驗證", "認證進度、測試週期、失效率改善", "可小幅計入基準情境，但要保留延遲風險。"],
            ["小量生產", "重複訂單、交期、產線投入", "檢查單位經濟與毛利，不只看營收。"],
            ["量產", "產能利用率、明確訂單、客戶擴廠", "納入基準EPS，並檢查客戶集中。"],
            ["平台／持續收入", "耗材、維護、軟體、升級與多產品導入", "可支持較高估值與較長持續性。"],
        ],
        [1300, 3900, 4160],
        first_col_bold=True,
    )

    doc.add_heading("6.2 必問的技術問題", level=2)
    for text in [
        "產品在機器人或自動化設備裡負責什麼？如果拿掉會發生什麼？",
        "最重要的規格是精度、扭矩密度、負載、速度、壽命、散熱、尺寸，還是控制整合？",
        "公司優勢來自設計、材料、加工、軟體、良率，還是客戶共同開發？",
        "量產時最難維持的是哪個規格？競爭者做出樣品是否等同能大量供貨？",
        "客戶重新設計或更換供應商，需要多久、多少成本與重新認證？",
        "產品平均壽命、維修頻率與替換週期如何？是否有耗材或服務收入？",
        "中國廠商的價格優勢有多大？品質差距正在縮小還是維持？",
        "美國／日本／歐洲設備商自行研發的動機與能力如何？哪些部分最可能垂直整合？",
        "公司有哪些專利、製程know-how、供應鏈或在地服務可以形成防禦？",
        "未來12個月最可能出現的可驗證事件是什麼？",
    ]:
        add_bullet(doc, text)

    doc.add_heading("6.3 中國低價與美國自研風險矩陣", level=2)
    add_table(doc,
        ["風險來源", "高風險特徵", "較能防禦的特徵", "要找的數據"],
        [
            ["中國低價競爭", "標準化零件、規格易複製、客戶只看價格。", "高可靠度、客製化、長認證週期、在地服務。", "ASP、毛利率、市占、退貨率、交期。"],
            ["美國／日本自研", "零件價值高、戰略安全、核心控制能力。", "供應商成本更低、開發速度快、跨客戶規模經濟。", "客戶研發投入、供應商名單、共同設計紀錄。"],
            ["技術替代", "單一機構方案、封閉規格、無軟體整合。", "平台化、多軸控制、軟硬整合、可升級。", "新產品路線、R&D占比、平台相容性。"],
        ],
        [1500, 2550, 2910, 2400],
        font_size=8.8,
        first_col_bold=True,
    )

    doc.add_heading("6.4 技術評分表（1-5分）", level=2)
    add_table(doc,
        ["指標", "權重", "1分", "3分", "5分"],
        [
            ["產品成熟度", "20%", "展示／送樣", "小量出貨", "已量產且重複訂單"],
            ["性能與可靠度", "20%", "無可驗證差異", "部分規格領先", "客戶流程高度依賴"],
            ["切換成本", "15%", "標準品隨時替換", "需重新調校", "需長期重新認證／設計"],
            ["量產與良率", "15%", "未知或不穩定", "可供貨但擴產待驗證", "規模量產紀錄完整"],
            ["價格與競爭防禦", "15%", "容易被低價取代", "性能可抵部分價差", "總持有成本明顯更低"],
            ["平台／持續收入", "15%", "一次性零件", "可交叉銷售", "耗材、維護、軟體或生態系"],
        ],
        [2000, 900, 2150, 2150, 2160],
        font_size=8.8,
        first_col_bold=True,
    )
    add_source_note(doc, "總分只用於比較與暴露分歧，不應把主觀分數當成精確機率。每項評分旁必須附一條來源或待驗證事項。")


def section_7(doc):
    doc.add_heading("7. 財務預測與一年目標價模型", level=1)
    doc.add_heading("7.1 從營運假設到EPS", level=2)
    finance_num_id = new_number_sequence(doc)
    for text in [
        "預估營收 = 既有產品營收 + 新產品出貨量 × 單價 + 其他成長／衰退。",
        "毛利 = 營收 × 毛利率；毛利率假設要連到產品組合、稼動率、價格與原料成本。",
        "營業利益 = 毛利 - 研發、銷售與管理費用。",
        "稅後淨利 = 營業利益 + 業外損益 - 稅負。",
        "EPS = 歸屬母公司稅後淨利 ÷ 稀釋後平均股數。",
        "一年目標價 = 未來12個月EPS × 一年後合理本益比。",
    ]:
        add_number(doc, text, finance_num_id)

    doc.add_heading("7.2 為何一定要做三種情境", level=2)
    add_table(doc,
        ["情境", "營運假設", "估值假設", "用途"],
        [
            ["悲觀", "訂單延遲、毛利下降、費用或稀釋高於預期。", "使用歷史較低或成熟同業倍數。", "判斷真實下檔是否可能超過-30%。"],
            ["基準", "現有訂單與產能正常轉化，不假設所有新題材成功。", "使用歷史中位數或合理同業倍數。", "主要決策目標價。"],
            ["樂觀", "新產品量產、客戶擴大、毛利與規模效益超預期。", "允許成長溢價，但不能無上限。", "理解上行選擇權和催化劑。"],
        ],
        [1200, 3650, 3000, 1510],
        first_col_bold=True,
    )
    add_para(doc, "建議第一版可先使用25%悲觀、50%基準、25%樂觀的機率權重。當證據增加後再調整，而不是為了得到理想目標價反向修改機率。")

    doc.add_heading("7.3 反推目標價：先問股價要求公司做到什麼", level=2)
    add_callout(doc, "公式", "要求的未來EPS = 目前股價 ×（1 + 期望報酬率）÷ 一年後本益比。", fill=GREEN_LIGHT, accent=GREEN)

    doc.add_heading("7.4 大銀微系統示範", level=2)
    add_para(doc, "截至2026年上半年，大銀合併營收18.79億元、毛利率37.97%、營益率17.35%、EPS 1.92元；2026年8月24日收盤價約198元。[S6][S7]")
    add_para(doc, "若要求一年上漲30%，目標價約為257.4元。不同一年後本益比，代表公司必須達到不同EPS：")
    add_table(doc,
        ["一年後本益比", "達成257.4元所需EPS", "相對上半年年化EPS 3.84元"],
        [
            ["35倍", "7.35元", "約需再提高91%"],
            ["40倍", "6.44元", "約需再提高68%"],
            ["45倍", "5.72元", "約需再提高49%"],
            ["50倍", "5.15元", "約需再提高34%"],
        ],
        [2200, 3100, 4060],
        aligns=[WD_ALIGN_PARAGRAPH.CENTER] * 3,
        first_col_bold=True,
    )
    add_callout(doc, "模型真正要回答的問題", "擴產、定位平台、半導體／矽光子設備與自動化元件，能否讓未來12個月EPS達到約5.2-7.4元？若成長未達預期，市場還願意給35-50倍本益比嗎？", fill=GOLD_LIGHT, accent=GOLD)

    doc.add_heading("7.5 波若威的反推提醒", level=2)
    add_para(doc, "波若威2026年上半年EPS約3.50元，2026年8月24日收盤約727元。[S8] 若同樣要求30%報酬，目標約945元；若一年後合理本益比為35倍，所需EPS約27元。這不是一般線性成長，而是對CPO／高速光通訊獲利跳升的高度期待。")
    add_para(doc, "一份2026年3月券商報告曾以2027年EPS 25.15元和36倍本益比估算905元目標價。[S14] 這可作為市場假設參考，但不是事實。研究重點應是量產節奏、良率、客戶與毛利是否逐季驗證。")

    doc.add_heading("7.6 估值倍數怎麼選", level=2)
    for text in [
        "先看公司過去3-5年本益比區間，但排除虧損或極端題材期間。",
        "再看商業模式相近的同業，而不是只因同屬機器人題材就比較。",
        "成長率、毛利率、現金流、客戶集中與持續性收入較好，才有理由給溢價。",
        "一年後成長若預計放緩，使用的退出本益比通常應低於目前倍數。",
        "若目標價必須同時假設EPS超預期和本益比不壓縮，模型安全邊際不足。",
    ]:
        add_bullet(doc, text)


def section_8(doc):
    doc.add_heading("8. 風險控制、停損與部位管理", level=1)
    add_callout(doc, "重要區分", "-30%是價格紀律，不是公司風險的上限。財報跳空、流動性不足或重大事件，都可能讓實際損失超過停損價。", fill=RED_LIGHT, accent=RED, text_color=RED)

    doc.add_heading("8.1 部位大小公式", level=2)
    add_callout(doc, "公式", "單一投資對總資產的最大損失 ≈ 部位占比 × 30%。例如持股占組合10%，跌30%時，組合約損失3%。")
    add_para(doc, "開會時還要決定『每個想法最多讓總資產虧多少』。假設容許單一想法使組合損失2%，使用30%停損時，部位上限約為2% ÷ 30% = 6.7%。")

    doc.add_heading("8.2 三種退出規則", level=2)
    add_table(doc,
        ["類型", "範例", "動作"],
        [
            ["價格停損", "自買進價下跌約30%。", "依事先規則退出或至少降至觀察部位，不臨時延後。"],
            ["基本面停損", "客戶驗證延後兩季、訂單取消、毛利無解惡化、重大稀釋。", "不等待價格到-30%，立即重估或退出。"],
            ["時間停損", "原定一年內催化劑未出現，故事持續往後延。", "在每季財報後檢查；不能用更遠的故事永久延長持有。"],
        ],
        [1500, 4850, 3010],
        first_col_bold=True,
    )

    doc.add_heading("8.3 預先設定的基本面警報", level=2)
    for text in [
        "機器人／新產品只有展覽消息，兩季內仍無客戶、訂單或收入證據。",
        "營收增長但毛利率連續惡化，顯示價格競爭或產品組合不如預期。",
        "大客戶延後資本支出，或單一客戶／單一應用占比升高。",
        "擴產後稼動率不足、庫存與應收帳款明顯快於營收成長。",
        "公司用增資、可轉債或高額資本支出支持故事，但自由現金流未改善。",
        "中國競爭者通過同等認證，且價格差足以迫使公司降價。",
    ]:
        add_bullet(doc, text)


def section_9(doc):
    doc.add_heading("9. 公司研究優先順序", level=1)
    add_para(doc, "以下是研究排序，不是買進排名。排序依據是：能否學到完整方法、是否有一年內可驗證數據，以及是否代表不同風險類型。")
    add_table(doc,
        ["優先", "公司", "研究理由", "本輪最重要問題", "暫定角色"],
        [
            ["1", "大銀", "已有獲利加速、產品與機器人運動鏈直接相關，適合建立第一個完整模型。", "EPS成長可否支撐高估值？成長主要來自半導體還是機器人？", "台股核心候選"],
            ["2", "直得", "同樣有精密運動與機械手臂布局，可和大銀做技術、商業化及估值對照。", "微型手臂何時形成可辨識營收？", "台股比較候選"],
            ["3", "ISRG", "成熟手術機器人商業模式，有裝機、手術量與耗材／服務的持續性。", "品質是否已被估值充分反映？", "美股品質基準"],
            ["4", "SYM", "高成長倉儲機器人，符合高波動偏好，但客戶集中與執行風險大。", "部署成長能否轉為穩定毛利與客戶多元化？", "高風險候選"],
            ["5", "TER", "持有UR與MiR機器人資產，但公司主要獲利亦受AI半導體測試影響。", "是否接受『非純機器人』但獲利來源較多元？", "替代／比較候選"],
            ["另案", "波若威", "題材與財務預期強，但屬AI光通訊，分類和競爭者不同。", "CPO量產假設是否已被股價提前反映？", "AI光通訊專案"],
        ],
        [700, 1000, 3150, 3130, 1380],
        font_size=8.5,
        first_col_bold=True,
    )

    doc.add_heading("9.1 已知的美股營運基準", level=2)
    for text in [
        "ISRG：2026年第二季da Vinci與Ion合計手術量年增約16%，da Vinci裝機基礎達11,710台，年增約12%。[S9]",
        "TER：2026年第二季機器人營收約1億美元、年增33.4%，約占公司總營收7.5%；其餘主要來自半導體與產品測試。[S10]",
        "SYM：2026財年第三季營收約7.21億美元；但2025財年Walmart約占營收85%，客戶集中是核心風險。[S11][S12]",
        "ROK：營運分為智慧裝置、軟體與控制、生命週期服務，可作成熟工業自動化商業模式比較。[S13]",
    ]:
        add_bullet(doc, text)


def section_10(doc):
    doc.add_heading("10. 四週研究計畫與交付物", level=1)
    add_table(doc,
        ["階段", "主要工作", "朋友交付", "你交付", "共同決策"],
        [
            ["會前", "閱讀企劃書、各自選最看好／最懷疑公司。", "產品位置草圖", "初步財務快照", "確認會議問題"],
            ["第0週", "第一次90分鐘會議。", "技術問題清單", "研究母體與任務表", "選4-5家公司"],
            ["第1週", "產品、成熟度、競爭與客戶驗證。", "直得vs大銀技術比較", "公司財務共同比較", "淘汰至少1家"],
            ["第2週", "營收、毛利、EPS與現金流模型。", "營運假設可行性審查", "三情境估值模型", "選3家深度研究"],
            ["第3週", "同業、本益比、反推EPS與風險。", "替代風險與技術否決條件", "估值區間與部位建議", "形成暫定排名"],
            ["第4週", "模擬投資委員會。", "技術辯護／反方", "投資論點／反方", "買進、等待或淘汰"],
        ],
        [1100, 2650, 2050, 2050, 1510],
        font_size=8.7,
        first_col_bold=True,
    )

    doc.add_heading("10.1 每家公司標準交付格式", level=2)
    for text in [
        "一頁公司研究卡：公司如何賺錢、機器人關聯、技術優勢、最大疑點。",
        "標準表件：三年財務、三情境估值、六類風險、正反論點、催化劑與否決條件。",
    ]:
        add_bullet(doc, text)


def section_11(doc):
    heading = doc.add_heading("11. 最終選股與投資決策規則", level=1)
    heading.paragraph_format.page_break_before = True
    doc.add_heading("11.1 最終評分（每項1-5分）", level=2)
    add_table(doc,
        ["面向", "權重", "核心問題"],
        [
            ["一年內財務催化劑", "20%", "未來四季有什麼能明確推動營收與EPS？"],
            ["技術與客戶黏著", "20%", "客戶為何不能改買中國低價或自行研發？"],
            ["獲利與現金流品質", "15%", "成長是否伴隨毛利、現金流與股本品質？"],
            ["估值與安全邊際", "20%", "基準情境是否仍有25%-30%上漲空間？"],
            ["風險可監控性", "15%", "失敗是否能在一至兩季內被看見？"],
            ["管理層與資訊可信度", "10%", "公司過去說法是否能被後續數據驗證？"],
        ],
        [2500, 1000, 5860],
        first_col_bold=True,
    )

    doc.add_heading("11.2 進入買進討論的最低條件", level=2)
    for text in [
        "有一個不依賴所有遠期題材都成功的基準情境。",
        "基準目標價至少有25%-30%上漲空間；樂觀情境最好達60%以上。",
        "悲觀情境與基本面停損已寫清楚，且部位大小能使組合損失可承受。",
        "未來四季至少有兩個可以驗證的催化劑與明確檢查日期。",
        "投資論點不能同時依賴EPS大幅超預期和估值倍數維持歷史高檔。",
        "兩人都能說出最強反方理由；重大分歧被保留在會議紀錄中。",
    ]:
        add_bullet(doc, text)

    doc.add_heading("11.3 決策分類", level=2)
    add_table(doc,
        ["決策", "定義", "下一步"],
        [
            ["買進／分批", "基本面、估值和催化劑同時通過；部位與停損已設定。", "建立追蹤表，財報後重估。"],
            ["等待價格", "公司符合條件，但目前預期報酬不足。", "設定合理進場區間，不追逐短線題材。"],
            ["等待證據", "估值可能有空間，但量產、客戶或毛利仍未驗證。", "列出下一個可驗證事件與日期。"],
            ["淘汰", "分類錯誤、商業化太遠、風險不可控或估值需要完美情境。", "記錄淘汰理由，避免同一敘事反覆吸引。"],
        ],
        [1500, 5070, 2790],
        first_col_bold=True,
    )


def appendices(doc):
    add_page_break(doc)
    doc.add_heading("附錄A：單一公司研究卡（可直接複製）", level=1)
    fields = [
        ("公司／代號", ""),
        ("一句話說明如何賺錢", ""),
        ("機器人／自動化關聯", "直接／間接；目前營收證據："),
        ("現有獲利主力", ""),
        ("未來四季兩個催化劑", "1.　　　　　　　　　　　　　　2."),
        ("技術優勢", "性能／量產／切換成本／軟體整合："),
        ("主要競爭者", "中國：　　　　　美國：　　　　　日本／歐洲："),
        ("最大疑點", ""),
        ("悲觀／基準／樂觀EPS", ""),
        ("合理估值倍數", "悲觀：　　　基準：　　　樂觀："),
        ("目標價區間", ""),
        ("否決條件", "1.　　　　　　　　　　　　　　2."),
        ("下一個檢查日", ""),
    ]
    add_table(doc, ["欄位", "填寫內容"], fields, [2500, 6860], font_size=9.5, first_col_bold=True)

    doc.add_heading("附錄B：三情境估值表", level=1)
    add_table(doc,
        ["項目", "悲觀", "基準", "樂觀"],
        [
            ["未來12個月營收", "", "", ""],
            ["營收成長率", "", "", ""],
            ["毛利率", "", "", ""],
            ["營益率", "", "", ""],
            ["EPS", "", "", ""],
            ["一年後本益比", "", "", ""],
            ["目標價", "", "", ""],
            ["機率", "25%", "50%", "25%"],
            ["催化劑／風險", "", "", ""],
        ],
        [2700, 2220, 2220, 2220],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        first_col_bold=True,
    )
    add_para(doc, "機率加權目標價 = 悲觀目標價 × 悲觀機率 + 基準目標價 × 基準機率 + 樂觀目標價 × 樂觀機率。", bold=True, color=NAVY)

    doc.add_heading("附錄C：第一次會議紀錄", level=1)
    add_table(doc,
        ["紀錄項目", "內容"],
        [
            ["會議日期／參與者", ""],
            ["確定的研究母體", ""],
            ["移出同業比較的公司", ""],
            ["目前最看好的公司與理由", "你：　　　　　　　　　朋友："],
            ["目前最懷疑的公司與理由", "你：　　　　　　　　　朋友："],
            ["最大的三個未知問題", "1.　　　　　　　　　2.　　　　　　　　　3."],
            ["研究任務與負責人", ""],
            ["交付日期／格式", ""],
            ["下次會議日期與目標", ""],
            ["本次未解決的分歧", ""],
        ],
        [2600, 6760],
        font_size=9.5,
        first_col_bold=True,
    )


def sources(doc):
    add_page_break(doc)
    doc.add_heading("資料來源與使用限制", level=1)
    add_para(doc, "本文件中的即時財務與股價資料以2026年8月24日前後可取得資訊為基準。開會前應再次更新股價、最新月營收與財報。公司展望、媒體與券商預估都屬待驗證假設，不應視為保證。")

    source_items = [
        ("S1", "直得科技官方產品與公司網站", "https://chieftek.com/zh-hant/"),
        ("S2", "大銀微系統官方產品與投資人網站", "https://www.hiwinmikro.tw/zh"),
        ("S3", "波若威官方公司概況", "https://www.browave.com/investors/Category/20"),
        ("S4", "直得2026年上半年財報公告（EPS 1.30元）", "https://m.moneydj.com/f1a.aspx?a=78c10eec-f716-4fef-b000-24aa47d42395"),
        ("S5", "直得2026年8月24日市場與近四季財務快照", "https://info.ifa.ai/tw-stock/1597"),
        ("S6", "大銀微系統2026年上半年財報摘要", "https://ww2.money-link.com.tw/RealtimeNews/NewsContent.aspx?PU=0010&SN=2406735002"),
        ("S7", "大銀微系統2026年8月24日股價與最新營收", "https://tw.stock.yahoo.com/quote/4576.TW/"),
        ("S8", "波若威2026年上半年營收盈餘與2026年8月24日股價", "https://ww2.money-link.com.tw/TWStock/StockBasic.aspx?SymId=3163&TWMId=Basic_salm1"),
        ("S9", "Intuitive Surgical 2026年第二季財報新聞稿", "https://investor.intuitivesurgical.com/news-releases/news-release-details/intuitive-announces-second-quarter-earnings-6"),
        ("S10", "Teradyne 2026年第二季10-Q", "https://investors.teradyne.com/sec-filings/all-sec-filings/content/0001193125-26-327715/ter-20260628.htm"),
        ("S11", "Symbotic 2026財年第三季業績", "https://ir.symbotic.com/news-releases/news-release-details/symbotic-reports-third-quarter-fiscal-year-2026-results"),
        ("S12", "Symbotic 2025年10-K（Walmart客戶集中）", "https://www.sec.gov/Archives/edgar/data/1837240/000183724025000278/sym-20250927.htm"),
        ("S13", "Rockwell Automation 2025年10-K", "https://www.rockwellautomation.com/content/dam/rockwell-automation/documents/pdf/company/about-us/ir/2025/Q4-FY25-10K.pdf"),
        ("S14", "福邦證券2026年3月研究情境（波若威，僅供假設比較）", "https://www.gfortune.com.tw/Report/%E7%A6%8F%E9%82%A6%E8%82%A1%E5%B8%82%E6%97%A9%E5%A0%B1%2020260312.pdf"),
    ]
    for sid, label, url in source_items:
        p = doc.add_paragraph(style="Source Text")
        r = p.add_run(f"[{sid}] {label} - ")
        set_run_font(r, size=9, bold=True, color=BLACK)
        add_hyperlink(p, url, url)

    doc.add_heading("證據優先順序", level=2)
    add_para(doc, "公司財報／監管申報 > 法說會與公司產品資料 > 產業研究與客戶／競爭者資料 > 媒體報導 > 社群與市場傳聞。第三方預估可用來了解市場期待，但不能取代自己的情境模型。")

    add_callout(doc, "使用限制", "本文件用於研究方法與會議規劃，不構成個人化投資建議或報酬保證。高波動股票可能跳空，實際損失可能超過預設停損。", fill=RED_LIGHT, accent=RED, text_color=RED)


def build():
    doc = setup_document()
    add_cover(doc)
    add_toc_and_summary(doc)
    section_1(doc)
    section_2(doc)
    section_3(doc)
    section_4(doc)
    section_5(doc)
    section_6(doc)
    section_7(doc)
    section_8(doc)
    section_9(doc)
    section_10(doc)
    section_11(doc)
    appendices(doc)
    sources(doc)

    core = doc.core_properties
    core.title = "機器人與自動化股票研究企劃書"
    core.subject = "第一次研究會議、四週研究流程與一年目標價框架"
    core.author = "研究小組"
    core.keywords = "機器人, 自動化, 股票研究, 目標價, 會議企劃"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
